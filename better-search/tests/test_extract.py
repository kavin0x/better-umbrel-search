from pathlib import Path

from pypdf import PdfWriter

from better_search.chunk import pack_text
from better_search.classify import Decision
from better_search.extract import (
    extract_document,
    extract_html,
    extract_image,
    extract_pdf,
    extract_pptx,
    rtf_to_text,
    should_ocr,
)
from PIL import Image


def write_text_pdf(path: Path, text: str) -> None:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET\n".encode("latin-1")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
        f"4 0 obj\n<< /Length {len(content)} >>\nstream\n".encode("latin-1") + content + b"endstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    blob = b"%PDF-1.4\n"
    offsets = []
    for obj in objects:
        offsets.append(len(blob))
        blob += obj
    xref_at = len(blob)
    xref = [f"xref\n0 {len(objects) + 1}\n", "0000000000 65535 f \n"]
    xref.extend(f"{offset:010d} 00000 n \n" for offset in offsets)
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    path.write_bytes(blob + "".join(xref).encode("ascii") + trailer.encode("ascii"))


def test_pack_text_overlaps() -> None:
    chunks = pack_text("word " * 500, size=200, overlap=40)
    assert len(chunks) > 2
    assert chunks[0] != chunks[1]


def test_pdf_docx_and_pptx(tmp_path: Path) -> None:
    pdf = tmp_path / "invoice.pdf"
    write_text_pdf(pdf, "roof leak invoice")
    pieces = extract_pdf(pdf, lambda _path, _index: "")
    assert "roof leak" in pieces[0].text
    assert pieces[0].page == 1

    from docx import Document

    docx_path = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("quarterly roof repair")
    document.save(docx_path)
    docx_text = " ".join(piece.text for piece in extract_document(docx_path, Decision("index", "document")))
    assert "quarterly roof repair" in docx_text

    from pptx import Presentation

    pptx_path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "solar panel invoice"
    presentation.save(pptx_path)
    slides = extract_pptx(pptx_path)
    assert slides[0].page == 1
    assert "solar panel" in slides[0].text


def test_blank_pdf_uses_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)
    pieces = extract_pdf(path, lambda _path, _index: "scanned roof leak")
    assert pieces[0].text == "scanned roof leak"


def test_html_rtf_and_xlsx(tmp_path: Path) -> None:
    html = tmp_path / "page.html"
    html.write_text("<html><style>x</style><p>gutter quote</p></html>", encoding="utf-8")
    assert "gutter quote" in extract_html(html)[0].text
    assert "x" not in extract_html(html)[0].text
    assert "roof leak" in rtf_to_text(r"{\rtf1\ansi roof leak\par}")

    from openpyxl import Workbook

    sheet_path = tmp_path / "costs.xlsx"
    book = Workbook()
    book.active["A1"] = "copper gutter quote"
    book.save(sheet_path)
    text = " ".join(piece.text for piece in extract_document(sheet_path, Decision("index", "document")))
    assert "copper gutter quote" in text


def test_camera_photo_skips_ocr_and_screenshot_does_not(tmp_path: Path) -> None:
    camera = tmp_path / "camera.jpg"
    image = Image.new("RGB", (8, 8), (10, 20, 30))
    exif = Image.Exif()
    exif[271] = "Canon"
    image.save(camera, exif=exif)
    with Image.open(camera) as opened:
        assert should_ocr(opened) is False
    calls: list[int] = []
    info = extract_image(camera, ocr_enabled=True, ocr_fn=lambda _image: calls.append(1) or "nope")
    assert calls == []
    assert info.text == ""

    shot = tmp_path / "shot.png"
    Image.new("RGB", (8, 8), (255, 255, 255)).save(shot)
    info = extract_image(shot, ocr_enabled=True, ocr_fn=lambda _image: "whiteboard notes")
    assert info.text == "whiteboard notes"
