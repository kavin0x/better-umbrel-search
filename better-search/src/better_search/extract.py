from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from PIL import Image

from better_search.classify import Decision

logger = logging.getLogger(__name__)

MIN_PAGE_CHARS = 20

try:
    import pillow_heif

    pillow_heif.register_heif_opener()
except Exception:
    logger.debug("HEIC support unavailable", exc_info=True)


@dataclass(frozen=True)
class Piece:
    text: str
    page: int | None


@dataclass(frozen=True)
class ImageInfo:
    text: str
    captured_at: float | None


def should_ocr(image: Image.Image) -> bool:
    exif = image.getexif()
    make = exif.get(271) or exif.get(272)
    return not bool(make)


def exif_timestamp(image: Image.Image) -> float | None:
    exif = image.getexif()
    raw = exif.get(36867) or exif.get(306)
    if not raw:
        return None
    try:
        return datetime.strptime(str(raw), "%Y:%m:%d %H:%M:%S").timestamp()
    except ValueError:
        return None


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def rtf_to_text(data: str) -> str:
    without_hex = re.sub(r"\\'[0-9a-fA-F]{2}", " ", data)
    without_controls = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", without_hex)
    plain = without_controls.replace("{", " ").replace("}", " ")
    return " ".join(plain.split())


def _xml_text(payload: bytes) -> str:
    try:
        root = ElementTree.fromstring(payload)
    except ElementTree.ParseError:
        return ""
    return " ".join(part.strip() for part in root.itertext() if part and part.strip())


def extract_pdf(path: Path, ocr_page) -> list[Piece]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pieces: list[Piece] = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if len(text.strip()) < MIN_PAGE_CHARS:
            try:
                ocr_text = ocr_page(path, index) or ""
            except Exception:
                logger.debug("ocr failed for %s page %s", path, index, exc_info=True)
                ocr_text = ""
            if ocr_text.strip():
                text = ocr_text
        pieces.append(Piece(text, index + 1))
    return pieces


def default_ocr_page(path: Path, page_index: int) -> str:
    import pypdfium2 as pdfium
    import pytesseract

    document = pdfium.PdfDocument(str(path))
    try:
        page = document[page_index]
        bitmap = page.render(scale=1.5)
        image = bitmap.to_pil()
        return pytesseract.image_to_string(image) or ""
    finally:
        document.close()


def extract_docx(path: Path) -> list[Piece]:
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" ".join(cells))
    return [Piece("\n".join(parts), 1)]


def extract_pptx(path: Path) -> list[Piece]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    pieces: list[Piece] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(notes)
        pieces.append(Piece("\n".join(parts), index))
    return pieces


def extract_xlsx(path: Path) -> list[Piece]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        pieces: list[Piece] = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" ".join(cells))
            pieces.append(Piece("\n".join(rows), index))
        return pieces
    finally:
        workbook.close()


def extract_html(path: Path) -> list[Piece]:
    soup = BeautifulSoup(_read_text(path), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return [Piece(soup.get_text(" ", strip=True), 1)]


def extract_odt(path: Path) -> list[Piece]:
    with zipfile.ZipFile(path) as archive:
        payload = archive.read("content.xml")
    return [Piece(_xml_text(payload), 1)]


def extract_epub(path: Path) -> list[Piece]:
    texts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.lower().endswith((".html", ".xhtml", ".htm")):
                texts.append(_xml_text(archive.read(name)))
    return [Piece("\n".join(part for part in texts if part), 1)]


def extract_image(path: Path, *, ocr_enabled: bool, ocr_fn) -> ImageInfo:
    with Image.open(path) as image:
        captured = exif_timestamp(image)
        text = ""
        if ocr_enabled and should_ocr(image):
            frame = image.convert("RGB")
            try:
                text = ocr_fn(frame) or ""
            except Exception:
                logger.debug("image ocr failed for %s", path, exc_info=True)
                text = ""
        return ImageInfo(text=text, captured_at=captured)


def default_ocr_image(image: Image.Image) -> str:
    import pytesseract

    return pytesseract.image_to_string(image) or ""


def extract_document(path: Path, decision: Decision, ocr_page=default_ocr_page) -> list[Piece]:
    suffix = path.suffix.casefold()
    mime = decision.mime or ""
    if suffix == ".pdf" or mime == "application/pdf":
        return extract_pdf(path, ocr_page)
    if suffix == ".docx" or "wordprocessingml" in mime:
        return extract_docx(path)
    if suffix == ".pptx" or "presentationml" in mime:
        return extract_pptx(path)
    if suffix == ".xlsx" or "spreadsheetml" in mime:
        return extract_xlsx(path)
    if suffix in {".html", ".htm"} or mime == "text/html":
        return extract_html(path)
    if suffix == ".odt" or mime == "application/vnd.oasis.opendocument.text":
        return extract_odt(path)
    if suffix == ".epub" or mime == "application/epub+zip":
        return extract_epub(path)
    if suffix == ".rtf" or mime in {"application/rtf", "text/rtf"}:
        return [Piece(rtf_to_text(_read_text(path)), 1)]
    return [Piece(_read_text(path), 1)]
